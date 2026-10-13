import pytest
from playwright.sync_api import sync_playwright
from pathlib import Path
from openpyxl import load_workbook


@pytest.fixture
def page():

    with sync_playwright() as playwright:

        browser = playwright.chromium.launch(
            channel="msedge",
            headless=False
        )

        page = browser.new_page()

        page.goto(
            "https://qooxdoo.org/qxl.widgetbrowser/",
            wait_until="domcontentloaded"
        )

        yield page

        browser.close()

@pytest.fixture
def excel_data():
    file_path = Path(__file__).parent / "test_data.xlsx"

    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active

    headers = [cell.value for cell in sheet[1]]
    rows = []

    for values in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in values):
            continue

        rows.append(dict(zip(headers, values)))

    workbook.close()

    print("Excel data loaded:", rows)
    return rows